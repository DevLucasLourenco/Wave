import docx
import openpyxl

class Transmitter:
    "Scrolls through the provided files and find the keys. Sequentially creates a .xlsx file with those keys as columns"
    
    def __init__(self, docFiles, delimiter):
        self.docFiles = docFiles
        self.__delimiter = delimiter
        self.keys = list()
        
        self.extractKeys()
        
        
    def extractKeys(self):
        for docFile in self.docFiles:
            doc = docx.Document(docFile)
            for para in doc.paragraphs:
                text = para.text 
                startPos = 0
                
                while True:
                    startPos = text.find(self.__delimiter, startPos)
                    if startPos == -1:
                        break
                    endPos = text.find(self.__delimiter, startPos + len(self.__delimiter))
                    if endPos == -1:
                        break
                    key = text[startPos + len(self.__delimiter):endPos]
                    self.keys.append(key) if key not in self.keys else ...
                    startPos = endPos + len(self.__delimiter)
                
                
    def export(self, outputFile):
        if self.keys:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(self.keys)
            workbook.save(outputFile)
        
